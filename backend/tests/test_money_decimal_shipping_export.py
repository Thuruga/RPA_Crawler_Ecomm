from __future__ import annotations

from decimal import Decimal
from io import BytesIO
from typing import Any

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from api.routes_search import (
    CrossMarketplaceExportRequest,
    ExportItem,
    _sanitize_cell,
    export_cross_marketplace,
)
from core.models import SearchProductResult, ShippingInfo
from services.phase10_acceptance import (
    FreightMatrixRow,
    validate_freight_gate_results,
)


def _money(value: Any) -> Decimal:
    return Decimal(str(value))


async def _export_first_data_row(item: ExportItem) -> list[Any]:
    response = await export_cross_marketplace(
        CrossMarketplaceExportRequest(
            items=[item],
            search_query="camiseta masculina",
            target_sku="SKU-1",
        )
    )
    chunks: list[bytes] = []
    async for chunk in response.body_iterator:
        chunks.append(chunk if isinstance(chunk, bytes) else bytes(chunk))
    workbook = load_workbook(BytesIO(b"".join(chunks)), data_only=False)
    worksheet = workbook["Busca SKU"]
    return [cell.value for cell in next(worksheet.iter_rows(min_row=2, max_row=2))]


def test_search_product_landed_price_uses_decimal_coerced_addition():
    product = SearchProductResult(
        brand="marketplace",
        product_name="Produto",
        url="https://example.com/produto",
        price_full=0.1,
        shipping_price=0.2,
    )

    assert _money(product.landed_price) == Decimal("0.3")


def test_shipping_info_only_marks_explicit_available_free_freight():
    available_zero = ShippingInfo(price=0, status="Disponivel")
    provider_marked_free = ShippingInfo(
        price=None,
        status="Gratis",
        is_free_shipping=True,
    )
    missing_freight = ShippingInfo(price=None, status="Disponivel")
    blocked_zero = ShippingInfo(price=0, status="Bloqueado (anti-bot)")

    assert available_zero.is_free_shipping is True
    assert _money(provider_marked_free.price) == Decimal("0")
    assert provider_marked_free.is_free_shipping is True
    assert missing_freight.is_free_shipping is False
    assert blocked_zero.is_free_shipping is False


@pytest.mark.asyncio
async def test_export_preserves_missing_freight_as_a_calcular_without_false_zero():
    row = await _export_first_data_row(
        ExportItem(
            marketplace="Mercado Livre",
            seller="Loja",
            title="Produto",
            price=129.90,
            shipping_price=None,
            landed_price=0,
            is_free_shipping=False,
            url="https://mercadolivre.com.br/produto",
        )
    )

    assert row[18] is None
    assert _money(row[19]) == Decimal("129.9")


@pytest.mark.parametrize("raw", ["=cmd", "\t=cmd", " \r\n+cmd", " \t@cmd"])
def test_sanitize_cell_quotes_formula_after_leading_whitespace(raw: str):
    assert _sanitize_cell(raw) == "'" + raw

@pytest.mark.asyncio
async def test_export_sanitizes_formula_strings_but_keeps_money_numeric():
    assert _sanitize_cell("=cmd") == "'=cmd"

    row = await _export_first_data_row(
        ExportItem(
            marketplace="=SUM(1,1)",
            seller="+seller",
            title="@Produto",
            price=10.50,
            shipping_price=1.25,
            landed_price=11.75,
            is_free_shipping=False,
            url="-https://example.com/produto",
        )
    )

    assert row[0] == "'=SUM(1,1)"
    assert row[1] == "'+seller"
    assert row[4] == "'@Produto"
    assert row[2] == "'-https://example.com/produto"
    assert isinstance(row[12], (int, float))
    assert isinstance(row[18], (int, float))
    assert isinstance(row[19], (int, float))
    assert _money(row[12]) == Decimal("10.5")
    assert _money(row[18]) == Decimal("1.25")
    assert _money(row[19]) == Decimal("11.75")


@pytest.mark.parametrize("field_name", ["price", "shipping_price", "landed_price"])
def test_export_item_rejects_invalid_money_values(field_name: str):
    payload = {
        "marketplace": "amazon",
        "seller": "seller",
        "title": "Produto",
        "price": 10.0,
        "shipping_price": 1.0,
        "landed_price": 11.0,
        "url": "https://amazon.com.br/produto",
    }
    payload[field_name] = "NaN"

    with pytest.raises(ValidationError):
        ExportItem.model_validate(payload)


def test_phase10_freight_gate_keeps_failure_reasons_distinct_for_plan_10_03():
    rows = [
        FreightMatrixRow(
            integration="vtex",
            state="blocked",
            product_url="https://aramis.com.br/produto",
            zipcode="01310100",
            shipping_price="12.50",
            evidence={"provider": "vtex"},
        ),
        FreightMatrixRow(
            integration="shopify",
            state="stale",
            product_url="https://www.ricardoalmeida.com.br/produto",
            zipcode="01310100",
            shipping_price="9.90",
            evidence={"provider": "shopify"},
        ),
        FreightMatrixRow(
            integration="wake",
            state="temporary_failure",
            product_url="https://www.richards.com.br/produto",
            zipcode="01310100",
            shipping_price="9.90",
            evidence={"provider": "wake"},
        ),
        FreightMatrixRow(
            integration="amazon",
            state="unsupported",
            product_url="https://amazon.com.br/produto",
            zipcode="01310100",
            shipping_price="9.90",
            evidence={"provider": "amazon"},
        ),
        FreightMatrixRow(
            integration="netshoes",
            state="available",
            product_url="https://netshoes.com.br/produto",
            zipcode="01310100",
            shipping_price="invalid-money",
            evidence={"provider": "netshoes"},
        ),
    ]

    result = validate_freight_gate_results(rows)
    reasons = {(rejection.integration, rejection.reason) for rejection in result.rejections}

    assert result.passed is False
    assert ("vtex", "blocked_state") in reasons
    assert ("shopify", "stale_state") in reasons
    assert ("wake", "temporary_failure_state") in reasons
    assert ("amazon", "unsupported_supported_integration") in reasons
    assert ("netshoes", "invalid_money") in reasons
