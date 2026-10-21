import io
import pytest
from fastapi.testclient import TestClient
import openpyxl

from app import app


@pytest.fixture
def client():
    return TestClient(app)


def test_scrape_category_export_xlsx(client):
    payload = {
        "items": [
            {
                "product_name": "Camisa Polo Piquet Aramis",
                "brand": "Aramis",
                "price": 199.90,
                "price_original": 299.90,
                "url": "https://www.aramis.com.br/polo-piquet",
                "image_url": "https://img.aramis.com.br/polo.jpg",
                "sku": "POLO-001",
                "is_available": True,
            },
            {
                "product_name": "Camiseta Básica",
                "brand": "Aramis",
                "price": 89.90,
                "price_original": None,
                "url": "https://www.aramis.com.br/camiseta",
                "image_url": None,
                "sku": "CAM-002",
                "is_available": False,
            },
        ],
        "filename": "meu_relatorio.xlsx",
    }

    response = client.post(
        "/scrape-category/export",
        json=payload,
        headers={"X-API-Key": "dev-api-key"},
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment; filename=\"meu_relatorio.xlsx\"" in response.headers["content-disposition"]

    # Verify workbook content
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    assert "Produtos Extraídos" in wb.sheetnames
    ws = wb["Produtos Extraídos"]

    # Check header
    headers = [cell.value for cell in ws[1]]
    assert headers == [
        "Marca", "Vendedor", "URL do Produto", "Código do Produto",
        "Nome do Produto", "Cor", "Descrição do Produto", "Composição",
        "Cores Disponíveis", "Tamanhos Disponíveis", "Categoria", "Avaliação",
        "Preço Atual", "Preço com Desconto", "Valor do Desconto",
        "Frete - Transportadora", "Tempo de Frete", "Frete Grátis",
        "Valor do Frete", "Preço Total",
    ]

    # Check rows
    row1 = [cell.value for cell in ws[2]]
    assert row1[0] == "Aramis"
    assert row1[2] == "https://www.aramis.com.br/polo-piquet"
    assert row1[3] == "POLO-001"
    assert row1[4] == "Camisa Polo Piquet Aramis"
    assert row1[12] == 199.90
    assert row1[13] == 199.90
    assert row1[14] == 100

    row2 = [cell.value for cell in ws[3]]
    assert row2[0] == "Aramis"
    assert row2[4] == "Camiseta Básica"
    assert row2[12] == 89.90
